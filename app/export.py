from datetime import datetime
from io import BytesIO
from typing import Annotated

from docx import Document
from docx.shared import Cm, Pt
from pydantic import BaseModel, ConfigDict, Field, HttpUrl, StringConstraints

Text = Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=6000)]
Identifier = Annotated[str, StringConstraints(pattern=r"^[A-Za-z0-9_-]{1,128}$")]


class SourceReference(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    document_id: Identifier
    page: int | None = Field(default=None, ge=1)


class RetrievedSource(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    document_id: Identifier
    title: Text
    url: HttpUrl


class OpinionSummary(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    topic: Text
    haltung: Text
    offene_fragen: list[Text] = Field(default_factory=list, max_length=30)
    gegenpositionen: list[Text] = Field(default_factory=list, max_length=30)
    sources: list[SourceReference] = Field(default_factory=list, max_length=50)


def render_summary(
    summary: OpinionSummary,
    *,
    sources: dict[str, RetrievedSource],
    owner: str,
    created_at: datetime,
) -> bytes:
    for reference in summary.sources:
        if reference.document_id not in sources:
            raise ValueError(f"Unknown source: {reference.document_id}")
    if created_at.utcoffset() is None:
        raise ValueError("Timestamp must include a timezone")

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.core_properties.title = summary.topic
    document.core_properties.author = owner
    document.add_heading(summary.topic, level=0)
    document.add_paragraph("SYNTHETISCHE DEMODATEN - KEINE REALEN UNTERNEHMENSDATEN")
    document.add_paragraph(f"Erstellt fuer: {owner}\nZeitpunkt: {created_at.isoformat()}")
    document.add_heading("Haltung zum Thema", level=1)
    document.add_paragraph(summary.haltung)
    for title, entries in [
        ("Fragen, die noch gekl\u00e4rt werden m\u00fcssen", summary.offene_fragen),
        ("Gegenpositionen", summary.gegenpositionen),
    ]:
        document.add_heading(title, level=1)
        if entries:
            for entry in entries:
                document.add_paragraph(entry, style="List Bullet")
        else:
            document.add_paragraph("Keine Eintraege festgehalten.")
    document.add_heading("Quellen", level=1)
    if not summary.sources:
        document.add_paragraph("Keine Quellen angegeben.")
    for reference in summary.sources:
        source = sources[reference.document_id]
        page = f", Seite {reference.page}" if reference.page else ""
        document.add_paragraph(f"{source.document_id}: {source.title}{page}\n{source.url}")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
