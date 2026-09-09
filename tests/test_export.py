from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from app.export import OpinionSummary, RetrievedSource, SourceReference, render_summary


def test_word_round_trip_preserves_headings_content_and_verified_sources():
    source = RetrievedSource(
        document_id="decision-01",
        title="Fiktive Entscheidungsvorlage",
        url="https://demo.sharepoint.com/sites/demo/input/decision.pdf",
    )
    summary = OpinionSummary(
        topic="AlpenMobil Demo AG: Pilot",
        haltung="Ich bevorzuge den Pilot. <Keine Freigabe> & weitere Pruefung.",
        offene_fragen=["Wer verantwortet die Datenqualitaet?"],
        gegenpositionen=["Der Vollausbau koennte frueher Nutzen bringen."],
        sources=[SourceReference(document_id="decision-01", page=2)],
    )
    binary = render_summary(
        summary,
        sources={source.document_id: source},
        owner="Demo-Nutzer",
        created_at=datetime(2026, 9, 9, 12, 0, tzinfo=UTC),
    )
    document = Document(BytesIO(binary))
    headings = [
        paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"
    ]
    assert headings == [
        "Haltung zum Thema",
        "Fragen, die noch gekl\u00e4rt werden m\u00fcssen",
        "Gegenpositionen",
        "Quellen",
    ]
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for expected in [
        summary.topic,
        summary.haltung,
        summary.offene_fragen[0],
        summary.gegenpositionen[0],
        str(source.url),
        "decision-01",
        "Seite 2",
        "Demo-Nutzer",
        "2026-09-09T12:00:00+00:00",
        "SYNTHETISCHE DEMODATEN - KEINE REALEN UNTERNEHMENSDATEN",
    ]:
        assert expected in text


def test_empty_categories_are_present_and_undecided_stance_is_preserved():
    summary = OpinionSummary(topic="Pilot", haltung="Noch keine Haltung gebildet.")
    binary = render_summary(summary, sources={}, owner="Demo", created_at=datetime.now(UTC))
    text = "\n".join(paragraph.text for paragraph in Document(BytesIO(binary)).paragraphs)
    assert "Noch keine Haltung gebildet." in text
    assert text.count("Keine Eintraege festgehalten.") == 2
    assert "Keine Quellen angegeben." in text


def test_unretrieved_source_is_rejected_before_rendering():
    summary = OpinionSummary(
        topic="Pilot",
        haltung="Unentschieden",
        sources=[SourceReference(document_id="invented-document", page=1)],
    )
    with pytest.raises(ValueError, match="Unknown source"):
        render_summary(summary, sources={}, owner="Demo", created_at=datetime.now(UTC))


@pytest.mark.parametrize("extra", [{"destination": "/other-folder"}, {"owner": "someone-else"}])
def test_model_cannot_supply_destination_or_owner(extra):
    with pytest.raises(ValidationError):
        OpinionSummary(topic="Pilot", haltung="Unentschieden", **extra)


def test_blank_topic_and_invalid_page_are_rejected():
    with pytest.raises(ValidationError):
        OpinionSummary(topic="   ", haltung="Unentschieden")
    with pytest.raises(ValidationError):
        SourceReference(document_id="decision-01", page=0)
